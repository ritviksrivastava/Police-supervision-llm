"""Build the complete supervision note (case facts + drafted conclusion, lapses, directions) as Word or PDF.

Word files are written with python-docx. PDFs are produced by converting that Word file with LibreOffice, because
LibreOffice shapes Devanagari conjuncts correctly (most PDF libraries do not).
Files are written to a private temporary folder and removed after REPORT_TTL seconds.
"""
import glob, os, re, shutil, subprocess, tempfile, time, uuid

REPORT_DIR = os.path.join(tempfile.gettempdir(), "sn_reports")
REPORT_TTL = 15 * 60
WORD_FONT = "Nirmala UI"           # present on Windows, where the Word file will usually be opened
PDF_FONT = "Lohit Devanagari"      # present on the server that renders the PDF
DRAFT_NOTE = "यह टिप्पणी मसौदा सहायक (AI) की सहायता से तैयार की गई है तथा पर्यवेक्षण पदाधिकारी द्वारा जाँची गई है।"

_HEAD = re.compile(r"^((?:\(\d+\)|\d{2}\.)?\s*[^:]{1,40}?\s*:-)\s*(.*)$", re.S)


def _cleanup():
    os.makedirs(REPORT_DIR, mode=0o700, exist_ok=True)
    for path in glob.glob(os.path.join(REPORT_DIR, "*")):
        try:
            if time.time() - os.path.getmtime(path) > REPORT_TTL:
                shutil.rmtree(path) if os.path.isdir(path) else os.remove(path)
        except OSError:
            pass


def _set_font(run, font, size=None, bold=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    # Word draws Devanagari with the "complex script" (cs) font, size and bold settings, so set those too.
    # Elements are inserted in the order the Word schema requires, or Word reports the file as damaged.
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)
    if bold is not None:
        run.bold = bold
        if bold:
            rpr.get_or_add_bCs()
    if size:
        run.font.size = Pt(size)
        rpr.find(qn("w:sz")).addnext(rpr.makeelement(qn("w:szCs"), {qn("w:val"): str(int(size * 2))}))


def build_docx(facts_text, conclusion, lapses, directions, supervisor, mark_draft, path, font=WORD_FONT):
    """facts_text is the case-facts text the model received (build_case_facts output)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2.2)

    def para(text="", bold=False, size=12, align=None, heading=None, space_after=6, color=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.25
        if align:
            p.alignment = align
        if heading:
            r = p.add_run(heading + " ")
            _set_font(r, font, size, True)
        if text:
            r = p.add_run(text)
            _set_font(r, font, size, bold)
            if color:
                r.font.color.rgb = color
        return p

    lines = [l for l in facts_text.split("\n") if l.strip()]
    para(lines[0], bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)    # पर्यवेक्षण टिप्पणी
    if len(lines) > 1:
        para(lines[1], bold=True, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=10)  # case line
    for line in lines[2:]:
        m = _HEAD.match(line)
        if m:
            para(m.group(2), heading=m.group(1), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        else:
            para(line, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    n = sum(1 for l in lines if re.match(r"^\(\d+\)", l))
    conc = [p.strip() for p in (conclusion or "").split("\n") if p.strip()]
    para(conc[0] if conc else "", heading=f"({n + 1}) निष्कर्ष :-", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for p in conc[1:]:
        para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para((lapses or "शून्य।").strip(), heading=f"({n + 2}) त्रुटि :-", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(heading=f"({n + 3}) निर्देश :-", space_after=4)
    for d in [d.strip() for d in (directions or "").split("\n") if d.strip()]:
        p = para(d, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4)
        p.paragraph_format.left_indent = Cm(0.8)

    para(space_after=24)
    if supervisor:
        para(supervisor, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    if mark_draft:
        para(DRAFT_NOTE, size=9, color=RGBColor(0x66, 0x66, 0x66), space_after=0)
    doc.save(path)
    return path


def make_report(fmt, facts_text, conclusion, lapses, directions, supervisor, mark_draft, case_no=""):
    """Return the path of a .docx or .pdf containing the full note."""
    _cleanup()
    work = os.path.join(REPORT_DIR, uuid.uuid4().hex)
    os.makedirs(work, mode=0o700)
    stem = "supervision_note_" + (re.sub(r"[^0-9A-Za-z]+", "-", case_no).strip("-") or "draft")
    if fmt == "docx":
        return build_docx(facts_text, conclusion, lapses, directions, supervisor, mark_draft,
                          os.path.join(work, stem + ".docx"))
    src = build_docx(facts_text, conclusion, lapses, directions, supervisor, mark_draft,
                     os.path.join(work, stem + ".docx"), font=PDF_FONT)
    profile = os.path.join(work, "lo_profile")          # own profile: conversions can run side by side
    subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                    "--outdir", work, src], check=True, timeout=120,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shutil.rmtree(profile, ignore_errors=True)
    os.remove(src)
    return os.path.join(work, stem + ".pdf")
